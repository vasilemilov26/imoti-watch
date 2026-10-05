"""HTML → clean text, document (PDF/DOCX) → text, link/block helpers."""
from __future__ import annotations

import io
import re
from urllib.parse import urljoin, urldefrag

from bs4 import BeautifulSoup, Tag

NOISE_RE = re.compile(r"(^|[-_ ])(nav|navbar|menu|footer|header|breadcrumb|sidebar|cookie|social|share|search|lang)([-_ ]|$)", re.I)
DOC_EXT = (".pdf", ".docx", ".doc", ".rtf", ".odt")


def soup_of(html: str) -> BeautifulSoup:
    return BeautifulSoup(html, "lxml")


def strip_noise(soup: BeautifulSoup) -> BeautifulSoup:
    for t in soup(["script", "style", "noscript", "nav", "header", "footer", "form", "iframe", "svg"]):
        t.decompose()
    for t in list(soup.find_all(True)):
        if not isinstance(t, Tag) or t.attrs is None:
            continue
        cls = " ".join(t.get("class", []) or []) + " " + (t.get("id") or "")
        if cls.strip() and NOISE_RE.search(cls) and t.name not in ("body", "html", "main", "article"):
            # don't kill huge containers that hold the real content
            if len(t.get_text(" ", strip=True)) < 4000:
                t.decompose()
    return soup


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()


def main_text(html: str) -> str:
    soup = strip_noise(soup_of(html))
    main = soup.find("main") or soup.find("article") or soup.body or soup
    return clean(main.get_text(" ", strip=True))


def abs_url(base: str, href: str) -> str | None:
    if not href:
        return None
    href = href.strip()
    if href.startswith(("javascript:", "mailto:", "tel:", "#")):
        return None
    return urldefrag(urljoin(base, href))[0]


def is_doc(url: str) -> bool:
    return url.lower().split("?")[0].endswith(DOC_EXT)


def doc_text(content: bytes, url: str, ctype: str = "") -> str:
    u = url.lower().split("?")[0]
    try:
        if u.endswith(".pdf") or "pdf" in ctype:
            from pdfminer.high_level import extract_text
            return clean(extract_text(io.BytesIO(content), maxpages=15))
        if u.endswith(".docx") or "officedocument.wordprocessing" in ctype:
            import docx
            d = docx.Document(io.BytesIO(content))
            parts = [p.text for p in d.paragraphs]
            for tbl in d.tables:
                for row in tbl.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return clean("\n".join(parts))
        if u.endswith((".doc", ".rtf")):
            # crude: pull readable Cyrillic/Latin runs out of legacy binary
            raw = content.decode("cp1251", "ignore")
            runs = re.findall(r"[А-Яа-яA-Za-z0-9 .,:;№\-\"„“()/%]{6,}", raw)
            return clean(" ".join(runs))[:20000]
    except Exception:
        return ""
    return ""


def block_for_link(a: Tag, max_chars: int = 1500) -> Tag:
    """Context of a link: its own row/list item/card – never a container holding other notices."""
    href = a.get("href")
    node = a
    for _ in range(8):
        parent = node.parent
        if parent is None or parent.name in ("body", "html", "main", "table", "tbody", "ul", "ol"):
            break
        others = {x.get("href") for x in parent.find_all("a", href=True)} - {href}
        if others and len(parent.get_text(" ", strip=True)) > 40:
            # parent already holds other links – stop unless we are still inside our own row
            if parent.name in ("tr", "li", "article"):
                node = parent
            break
        node = parent
        if node.name in ("tr", "li", "article"):
            break
    if len(node.get_text(" ", strip=True)) > max_chars * 2:
        return a
    return node
